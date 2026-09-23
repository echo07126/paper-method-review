from pathlib import Path

from docx import Document

from app.engine.checklist import load_checklist
from app.engine.reviewer import review_document
from app.parsers.registry import parse_document

ROOT = Path(__file__).resolve().parents[2]


def build_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("方法", level=1)
    paragraph = doc.add_paragraph()
    run = paragraph.add_run("该结论引用了一条不存在的文献")
    normal = paragraph.add_run("。")
    normal.font.superscript = False
    cite = paragraph.add_run("[99]")
    cite.font.superscript = True
    doc.add_heading("References", level=1)
    doc.add_paragraph("[1] Some reference. 2026.")
    doc.save(path)


def test_parser_detects_superscript_citation(tmp_path: Path) -> None:
    source = tmp_path / "sample.docx"
    build_docx(source)
    document = parse_document(source, "sample.docx", None)
    assert document.sections
    assert any(99 in citation.unresolved for citation in document.citations)


def test_rules_report_dangling_citation(tmp_path: Path) -> None:
    source = tmp_path / "sample.docx"
    build_docx(source)
    document = parse_document(source, "sample.docx", None)
    items = load_checklist(ROOT / "checklists" / "quant-ai-v1.json")
    report = review_document(document, items, provider=None, use_llm=False)
    assert report.counts["total"] >= 1
    assert all(finding.provenance.get("engine") == "rule" for finding in report.findings)
