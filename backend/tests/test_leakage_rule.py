from pathlib import Path

from docx import Document

from app.engine.elements import extract_elements
from app.engine.rules import run_rules
from app.parsers.registry import parse_document


def build(path: Path, paragraphs: list[str]) -> None:
    doc = Document()
    doc.add_heading("4. 实验", level=1)
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(path)


def verdicts(path: Path) -> dict[str, str]:
    document = parse_document(path, path.name, None)
    elements = extract_elements(document)
    return {f.checklist_item_id: f.verdict.value for f in run_rules(document, elements)}


def test_validation_set_tuning_is_not_leakage(tmp_path: Path) -> None:
    source = tmp_path / "proper.docx"
    build(source, [
        "D1 按类别分层划分为训练集、验证集与测试集，比例为 8:1:1，划分以随机种子 42 固定。",
        "验证集用于超参数选择与早停监控，测试集仅用于最终评估一次。",
    ])
    assert verdicts(source).get("R-03") == "pass"


def test_test_set_tuning_is_leakage(tmp_path: Path) -> None:
    source = tmp_path / "leak.docx"
    build(source, [
        "超参数（学习率、批大小）在测试集上通过网格搜索确定，并在同一测试集上评估最终性能。",
        "模型在训练集上训练，测试集用于调参与模型选择。",
    ])
    assert verdicts(source).get("R-03") == "problem"
