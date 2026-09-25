from pathlib import Path

from docx import Document

from app.engine.checklist import load_checklist
from app.engine.elements import extract_elements, has_element
from app.engine.figures import demote_for_figures, find_figure_references
from app.engine.reviewer import review_document
from app.models.schemas import Anchor, Finding, Severity, Verdict
from app.parsers.registry import parse_document

ROOT = Path(__file__).resolve().parents[2]


def test_figure_reference_detected_and_noted(tmp_path: Path) -> None:
    source = tmp_path / "figures.docx"
    doc = Document()
    doc.add_heading("2. Methods", level=1)
    doc.add_paragraph("We trained the model with a learning rate of 0.001.")
    doc.add_paragraph("As shown in Figure 1, the proposed model outperforms the baseline.")
    doc.add_paragraph("The detailed results are listed in Table 2.")
    doc.save(source)

    document = parse_document(source, source.name, None)
    references = find_figure_references(document)
    assert len(references) >= 2, "应识别图/表引用"

    report = review_document(document, load_checklist(ROOT / "checklists" / "quant-ai-v1.json"))
    assert any("图表引用" in note for note in report.notes), "报告应提示图表未解析"


def test_demote_for_figures_only_targets_evidence_items() -> None:
    findings = [
        Finding(finding_id="F1", checklist_item_id="R-09", verdict=Verdict.PROBLEM, severity=Severity.MID, headline="缺少置信区间", provenance={"engine": "rule", "rule_id": "R-09"}),
        Finding(finding_id="F2", checklist_item_id="R-05", verdict=Verdict.PROBLEM, severity=Severity.MID, headline="缺少消融", provenance={"engine": "rule", "rule_id": "R-05"}),
        Finding(finding_id="F3", checklist_item_id="R-09", verdict=Verdict.PASS, severity=Severity.LOW, headline="通过", provenance={"engine": "rule", "rule_id": "R-09"}),
    ]
    demoted = {f.finding_id: f for f in demote_for_figures(findings)}
    assert demoted["F1"].verdict == Verdict.UNCERTAIN
    assert demoted["F1"].provenance.get("evidence_gate") == "figures_unparsed"
    assert demoted["F2"].verdict == Verdict.PROBLEM, "消融类结论不因图表降级"
    assert demoted["F3"].verdict == Verdict.PASS

def test_demote_flag_downgrades_figure_dependent_findings(tmp_path: Path) -> None:
    source = tmp_path / "figure_only_stats.docx"
    doc = Document()
    doc.add_heading("3. Results", level=1)
    doc.add_paragraph("Figure 1 shows the accuracy of 0.91 for the proposed model.")
    doc.save(source)

    document = parse_document(source, source.name, None)
    items = load_checklist(ROOT / "checklists" / "quant-ai-v1.json")

    strict = review_document(document, items, demote_on_figures=False)
    conservative = review_document(document, items, demote_on_figures=True)

    strict_r09 = [f for f in strict.findings if f.checklist_item_id == "R-09"][0]
    conservative_r09 = [f for f in conservative.findings if f.checklist_item_id == "R-09"][0]

    assert strict_r09.verdict == Verdict.PROBLEM
    assert conservative_r09.verdict == Verdict.UNCERTAIN
    assert conservative_r09.provenance.get("evidence_gate") == "figures_unparsed"
    assert conservative.figure_references >= 1
    assert conservative.demote_on_figures is True
    assert conservative.counts["total"] < strict.counts["total"]


def test_table_evidence_blocks_demotion(tmp_path: Path) -> None:
    """表体已被结构化并抽出该条目证据时，不应再按「图表未解析」降级（需求 15.5.2）。"""
    source = tmp_path / "table_evidence.docx"
    doc = Document()
    doc.add_heading("2. Methods", level=1)
    doc.add_paragraph("As shown in Figure 1, the cohort details are summarised below.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Cohort"
    table.cell(0, 1).text = "Size"
    table.cell(1, 0).text = "训练队列"
    table.cell(1, 1).text = "42 例"
    doc.save(source)

    document = parse_document(source, source.name, None)
    elements = extract_elements(document)
    assert document.tables and has_element(elements, "sample_size"), "表内样本量应被抽取"

    items = load_checklist(ROOT / "checklists" / "quant-ai-v1.json")
    strict = review_document(document, items, demote_on_figures=False)
    conservative = review_document(document, items, demote_on_figures=True)

    strict_r02 = [f for f in strict.findings if f.checklist_item_id == "R-02"][0]
    conservative_r02 = [f for f in conservative.findings if f.checklist_item_id == "R-02"][0]

    assert strict_r02.verdict == Verdict.PROBLEM
    assert conservative_r02.verdict == Verdict.PROBLEM, "表格已解析且命中证据，不应降级"
    assert conservative_r02.provenance.get("evidence_gate") != "figures_unparsed"