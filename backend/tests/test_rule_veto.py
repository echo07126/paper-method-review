import json
from pathlib import Path

from docx import Document

from app.engine.checklist import load_checklist
from app.engine.llm_provider import LLMResult
from app.engine.reviewer import review_document
from app.parsers.registry import parse_document

ROOT = Path(__file__).resolve().parents[2]


class VetoProvider:
    """规则已判定 R-07 通过（有 p 值且有 t 检验），模型却报问题 —— 应被否决。"""

    def complete_structured(self, messages, temperature: float = 0.0):
        if "逐字引用" in messages[0]["content"]:
            return LLMResult(content=json.dumps({"results": []}), model="veto-fake", tokens={"input": 1, "output": 1}, retries=0, duration_ms=1)
        content = json.dumps([
            {"checklist_item_id": "R-07", "verdict": "problem", "severity": "high", "headline": "缺少统计检验", "description": "模型误判", "suggestion": "补检验"},
        ])
        return LLMResult(content=content, model="veto-fake", tokens={"input": 1, "output": 1}, retries=0, duration_ms=1)


def build_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("3. 结果", level=1)
    doc.add_paragraph("组间差异采用配对 t 检验，P<0.05。")
    doc.save(path)


def test_rule_pass_vetoes_llm_problem(tmp_path: Path) -> None:
    source = tmp_path / "veto.docx"
    build_docx(source)
    document = parse_document(source, "veto.docx", None)
    items = load_checklist(ROOT / "checklists" / "quant-ai-v1.json")

    report = review_document(document, items, provider=VetoProvider(), use_llm=True)
    r07 = [f for f in report.findings if f.checklist_item_id == "R-07"][0]
    assert r07.verdict.value == "pass", "规则判定通过时应否决模型的 problem"
    assert r07.provenance.get("vetoed_llm") is True
    assert r07.finding_id not in {f.finding_id for f in report.findings if f.verdict.value == "problem"}
