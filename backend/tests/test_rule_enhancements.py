from pathlib import Path

from docx import Document

from app.engine.elements import extract_elements
from app.engine.rules import run_rules
from app.parsers.registry import parse_document


def verdicts(paragraphs: list[str], tmp_path: Path, name: str) -> dict[str, str]:
    path = tmp_path / name
    doc = Document()
    doc.add_heading("4. 实验", level=1)
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(path)
    document = parse_document(path, path.name, None)
    elements = extract_elements(document)
    return {f.checklist_item_id: f.verdict.value for f in run_rules(document, elements)}


def test_split_detail_required(tmp_path: Path) -> None:
    vague = verdicts([
        "模型在验证集上通过网格搜索调参，并在测试集上评估性能。",
        "该方法在多个数据集上表现良好。",
    ], tmp_path, "vague.docx")
    assert vague.get("R-03") == "problem", "只提验证集/测试集但无比例与种子应判问题"

    detailed = verdicts([
        "数据按类别分层划分，比例为 8:1:1：训练集 32000、验证集 4000、测试集 4000，随机种子固定为 42。",
        "测试集仅用于最终评估一次。",
    ], tmp_path, "detailed.docx")
    assert detailed.get("R-03") == "pass"


def test_single_baseline_flagged(tmp_path: Path) -> None:
    single = verdicts([
        "本文提出一种多尺度卷积模型，在测试集上准确率达到 0.96。",
        "为验证有效性，本文以 BERT 作为对比基线并微调。",
    ], tmp_path, "single.docx")
    assert single.get("R-04") == "problem", "仅 1 个基线应判问题"

    multiple = verdicts([
        "本文提出一种多尺度卷积模型，在测试集上准确率达到 0.96。",
        "对比方法为 4 个基线：TextCNN、TextRCNN、DPCNN 与 BERT。",
    ], tmp_path, "multiple.docx")
    assert multiple.get("R-04") == "pass"


def test_classification_metric_required(tmp_path: Path) -> None:
    engineering_only = verdicts([
        "本文提出多尺度卷积分类模型 MCAN，用于文本分类。",
        "本文报告参数量、训练时间与单条推理耗时三项工程指标。",
    ], tmp_path, "eng.docx")
    assert engineering_only.get("R-06") == "problem", "分类任务只报工程指标应判问题"

    with_metric = verdicts([
        "本文提出多尺度卷积分类模型 MCAN，用于文本分类。",
        "评价指标采用准确率与 macro-F1，测试集准确率为 0.962。",
    ], tmp_path, "metric.docx")
    assert with_metric.get("R-06") == "pass"


def test_named_test_required_for_p_value(tmp_path: Path) -> None:
    vague = verdicts([
        "对两组性能差异进行检验，结果表明差异具有统计学意义（p<0.05）。",
        "模型在训练集与测试集上的准确率分别为 0.95 与 0.93。",
    ], tmp_path, "vague_test.docx")
    assert vague.get("R-07") == "problem", "未说明检验方法应判问题"

    named = verdicts([
        "采用配对 t 检验比较两组性能差异（p<0.05）。",
        "模型在训练集与测试集上的准确率分别为 0.95 与 0.93。",
    ], tmp_path, "named_test.docx")
    assert named.get("R-07") == "pass"


def test_hyperparameter_range_required(tmp_path: Path) -> None:
    no_range = verdicts([
        "训练采用网格搜索确定超参数，学习率为 5e-4，批量大小为 128。",
        "模型在测试集上的准确率为 0.95。",
    ], tmp_path, "no_range.docx")
    assert no_range.get("R-12") == "problem", "未给搜索范围应判问题"

    with_range = verdicts([
        "超参数以网格搜索确定，学习率取值范围 [1e-4, 3e-3]，批量大小候选 {64, 128}。",
        "模型在测试集上的准确率为 0.95。",
    ], tmp_path, "range.docx")
    assert with_range.get("R-12") == "pass"
