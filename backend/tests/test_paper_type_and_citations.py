from pathlib import Path

from docx import Document

from app.engine.checklist import load_checklist
from app.engine.elements import extract_elements, has_element
from app.engine.paper_type import classify
from app.engine.reviewer import review_document
from app.engine.rules import run_rules
from app.parsers.registry import parse_document

ROOT = Path(__file__).resolve().parents[2]

REVIEW_TEXT = [
    "摘要：本文以机器学习为视角，系统阐述线性回归的模型设定与参数估计，并梳理研究进展。",
    "1 引言",
    "线性回归是机器学习中最基础的模型，广泛用于经济预测与科学计算。",
    "2 方法与研究进展",
    "设训练集由 N 个样本组成，最小二乘估计使残差平方和最小。",
    "岭回归在目标中加入 L2 惩罚，正则化系数取值范围 [1e-4, 1e-1]，通常通过交叉验证确定。",
    "3 结语",
    "本文的局限在于未覆盖核方法与贝叶斯扩展，未来可进一步补充。",
]


def build(path: Path, paragraphs: list[str]) -> None:
    doc = Document()
    for index, text in enumerate(paragraphs):
        if text[:2] in {"1 ", "2 ", "3 "}:
            doc.add_heading(text, level=1)
        else:
            doc.add_paragraph(text)
    doc.save(path)


def test_review_paper_skips_empirical_checks(tmp_path: Path) -> None:
    source = tmp_path / "review.docx"
    build(source, REVIEW_TEXT)
    document = parse_document(source, source.name, None)

    kind, _evidence = classify(document)
    assert kind == "review"

    items = load_checklist(ROOT / "checklists" / "quant-ai-v1.json")
    report = review_document(document, items)
    problems = {f.checklist_item_id for f in report.findings if f.verdict.value == "problem"}
    assert problems == set(), f"综述论文不应产生实证类问题，实际：{problems}"
    assert any("综述/理论论文" in note for note in report.notes)
    assert report.paper_type == "review"


def test_research_goal_and_hyperparameter_patterns(tmp_path: Path) -> None:
    source = tmp_path / "goal.docx"
    build(source, REVIEW_TEXT)
    document = parse_document(source, source.name, None)
    elements = extract_elements(document)
    assert has_element(elements, "research_goal"), "“本文以…系统阐述”应识别为研究目标"
    assert has_element(elements, "hyperparameter")

    items = load_checklist(ROOT / "checklists" / "quant-ai-v1.json")
    rule_verdicts = {f.checklist_item_id: f.verdict.value for f in run_rules(document, elements)}
    assert rule_verdicts.get("R-01") == "pass"
    assert rule_verdicts.get("R-12") == "pass", "“通过交叉验证确定”应视为超参选择依据"

    report = review_document(document, items)
    assert {f.checklist_item_id: f.verdict.value for f in report.findings}.get("R-01") == "pass"


def test_math_superscripts_are_not_citations(tmp_path: Path) -> None:
    source = tmp_path / "math.docx"
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("决定系数 R")
    sup = paragraph.add_run("2")
    sup.font.superscript = True
    paragraph.add_run(" 随特征数增加而增大，参见文献")
    citation = paragraph.add_run("[1]")
    citation.font.superscript = True
    paragraph.add_run("的讨论。")
    doc.save(source)

    document = parse_document(source, source.name, None)
    tokens = [citation.text for citation in document.citations]
    assert tokens == ["[1]"], f"只应识别引用样式角标，实际：{tokens}"
