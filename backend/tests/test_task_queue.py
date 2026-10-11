"""任务队列（重接口异步化）：入队 202、状态流转、跨会话隔离、失败落库、限流。

覆盖点对应《重接口异步任务化改造计划》§三：`POST` 立即 202 + 轮询 `GET /tasks/{id}`。
"""
import time
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pydantic_settings import SettingsConfigDict

from app.core import ratelimit
from app.core.config import Settings, set_settings_override
from app.core.errors import AppError
from app.main import app
from app.storage.db import init_db
from app.storage.repository import Store
from app.tasks import shutdown_executor, submit

ROOT = Path(__file__).resolve().parents[2]


class IsolatedSettings(Settings):
    """断开 backend/.env，使显式入参成为唯一配置来源。"""

    model_config = SettingsConfigDict(env_file=None, extra="ignore")


def _build_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("方法", level=1)
    doc.add_paragraph("本研究采用五折交叉验证，并与基线方法对比。")
    doc.add_heading("References", level=1)
    doc.add_paragraph("[1] Some reference. 2026.")
    doc.save(path)


@pytest.fixture()
def env(tmp_path: Path):
    settings = IsolatedSettings(
        db_path=str(tmp_path / "app.db"),
        temp_dir=str(tmp_path / "tmp"),
        checklist_dir=str(ROOT / "checklists"),
        app_env="dev",
        docs_enabled=False,
    )
    set_settings_override(settings)
    init_db(settings.db_path)
    ratelimit._WINDOWS.clear()
    try:
        with TestClient(app) as client:
            yield client, settings
    finally:
        set_settings_override(None)
        ratelimit._WINDOWS.clear()
        shutdown_executor()


def _upload(client: TestClient, tmp_path: Path) -> str:
    sample = tmp_path / "sample.docx"
    _build_docx(sample)
    with sample.open("rb") as handle:
        response = client.post(
            "/api/v1/uploads",
            files={"file": ("sample.docx", handle, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
    assert response.status_code == 200, response.text
    return response.json()["document_id"]


def _poll(client: TestClient, task_id: str, timeout: float = 10.0) -> dict:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        response = client.get(f"/api/v1/tasks/{task_id}")
        assert response.status_code == 200, response.text
        body = response.json()
        if body["state"] in {"succeeded", "failed"}:
            return body
        time.sleep(0.05)
    raise AssertionError("任务未在时限内完成")


def test_review_task_is_accepted_and_pollable(env, tmp_path: Path) -> None:
    """POST /reviews 立即 202 返回 task_id，轮询后 result 与旧同步响应字段一致。"""
    client, _ = env
    document_id = _upload(client, tmp_path)

    accepted = client.post("/api/v1/reviews", json={"document_id": document_id, "use_llm": False})
    assert accepted.status_code == 202, accepted.text
    body = accepted.json()
    assert body["kind"] == "review" and body["state"] == "queued"
    assert body["poll_url"] == f"/tasks/{body['task_id']}"

    done = _poll(client, body["task_id"])
    assert done["state"] == "succeeded", done
    assert done["kind"] == "review"
    result = done["result"]
    for field in ("report_id", "counts", "duration_ms", "tokens", "notes", "figure_references", "paper_type", "quality"):
        assert field in result, f"result 缺少旧响应字段 {field}"
    assert client.get(f"/api/v1/reports/{result['report_id']}").status_code == 200


def test_task_is_scoped_to_session(env, tmp_path: Path) -> None:
    """另一个会话（无 cookie）轮询同一 task_id 必须 404。"""
    client, _ = env
    document_id = _upload(client, tmp_path)
    accepted = client.post("/api/v1/reviews", json={"document_id": document_id, "use_llm": False})
    task_id = accepted.json()["task_id"]

    other = TestClient(app)
    assert other.get(f"/api/v1/tasks/{task_id}").status_code == 404


def test_task_reports_failure_without_raising(tmp_path: Path) -> None:
    """任务体抛错只把错误写进 job（不崩工作线程）：AppError 记 code/message，其它记 internal_error。"""
    settings = IsolatedSettings(
        db_path=str(tmp_path / "app.db"),
        temp_dir=str(tmp_path / "tmp"),
        checklist_dir=str(ROOT / "checklists"),
        app_env="dev",
    )
    set_settings_override(settings)
    init_db(settings.db_path)
    try:
        store = Store(settings.db_path)
        session_id, _ = store.create_session(2)

        def boom() -> dict:
            raise AppError("bad_request", "业务校验失败", 400)

        def crash() -> dict:
            raise ValueError("意外错误")

        failed = _wait(store, session_id, submit("review", session_id, {}, boom))
        assert failed["state"] == "failed"
        assert failed["error"] == {"code": "bad_request", "message": "业务校验失败"}

        crashed = _wait(store, session_id, submit("review", session_id, {}, crash))
        assert crashed["state"] == "failed"
        assert crashed["error"]["code"] == "internal_error"
    finally:
        set_settings_override(None)
        shutdown_executor()


def test_task_transitions_through_running(tmp_path: Path) -> None:
    """超时任务可观测到 running 中间态，最终 succeeded 且 result 落库。"""
    settings = IsolatedSettings(
        db_path=str(tmp_path / "app.db"),
        temp_dir=str(tmp_path / "tmp"),
        checklist_dir=str(ROOT / "checklists"),
        app_env="dev",
    )
    set_settings_override(settings)
    init_db(settings.db_path)
    try:
        store = Store(settings.db_path)
        session_id, _ = store.create_session(2)
        job_id = submit("review", session_id, {}, lambda: (time.sleep(0.2), {"ok": True})[1])

        observed_running = False
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            state = store.authorized_job(session_id, job_id)["state"]
            if state == "running":
                observed_running = True
            if state == "succeeded":
                break
            time.sleep(0.02)

        body = store.authorized_job(session_id, job_id)
        assert body["state"] == "succeeded"
        assert body["result"] == {"ok": True}
        assert observed_running, "应能观测到 running 中间态"
    finally:
        set_settings_override(None)
        shutdown_executor()


def test_eleventh_review_submission_is_rate_limited(env, tmp_path: Path) -> None:
    """限流仍按「提交时刻」扣桶：第 11 次审查提交返回 429。"""
    client, settings = env
    document_id = _upload(client, tmp_path)

    statuses = [
        client.post("/api/v1/reviews", json={"document_id": document_id, "use_llm": False}).status_code
        for _ in range(settings.rate_limit_reviews_per_minute + 1)
    ]
    assert statuses[: settings.rate_limit_reviews_per_minute] == [202] * settings.rate_limit_reviews_per_minute
    assert statuses[-1] == 429, statuses


def _wait(store: Store, session_id: str, job_id: str, timeout: float = 5.0) -> dict:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        body = store.authorized_job(session_id, job_id)
        if body["state"] in {"succeeded", "failed"}:
            return body
        time.sleep(0.02)
    raise AssertionError("任务未在时限内完成")