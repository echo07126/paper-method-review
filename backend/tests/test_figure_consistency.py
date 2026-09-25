from pathlib import Path

from docx import Document

from app.engine.figure_consistency import check_figure_consistency
from app.parsers.registry import parse_document


def _checks(document) -> set[str]:
    return {f.provenance.get("figure_check") for f in check_figure_consistency(document)}


def test_excerpt_without_figures_is_not_flagged(tmp_path: Path) -> None:
    """无图片、无表格的节选样本不做任何图表一致性判定，避免误报。"""
    source = tmp_path / "excerpt.docx"
    doc = Document()
    doc.add_heading("2. 方法", level=1)
    doc.add_paragraph("本文提出的系统流程如图1所示，包含检测与分类两个阶段。")
    doc.save(source)

    document = parse_document(source, source.name, None)
    assert check_figure_consistency(document) == []


def test_caption_missing_and_reference_without_caption(tmp_path: Path) -> None:
    """有图但无题注、且正文引用了图1 → 两条附加检查同时报出。"""
    source = tmp_path / "ref_only.docx"
    doc = Document()
    doc.add_heading("2. 方法", level=1)
    doc.add_paragraph("本文提出的系统流程如图1所示，包含检测与分类两个阶段。")
    doc.save(source)

    document = parse_document(source, source.name, None)
    document.meta["image_count"] = 1  # 该样本无内嵌图片，注入计数以模拟「有图但缺题注」

    checks = _checks(document)
    assert "caption_missing" in checks
    assert "reference_without_caption" in checks


def test_caption_and_reference_match_is_clean(tmp_path: Path) -> None:
    """题注与正文引用一一对应时不报任何问题。"""
    source = tmp_path / "matched.docx"
    doc = Document()
    doc.add_heading("3. 结果", level=1)
    doc.add_paragraph("如图1所示，模型准确率达到0.91。")
    doc.add_paragraph("图1 模型准确率对比")
    doc.add_paragraph("表1列出了主结果。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Model"
    table.cell(0, 1).text = "Accuracy"
    table.cell(1, 0).text = "Ours"
    table.cell(1, 1).text = "0.91"
    doc.add_paragraph("表1 主结果")
    doc.save(source)

    document = parse_document(source, source.name, None)
    document.meta["image_count"] = 1

    assert check_figure_consistency(document) == []


def test_matching_sample_size_is_not_flagged(tmp_path: Path) -> None:
    """题注与正文声明的样本量一致时不报冲突。"""
    source = tmp_path / "consistent.docx"
    doc = Document()
    doc.add_heading("3. 结果", level=1)
    doc.add_paragraph("如图1所示，本研究纳入42例受试者的准确率结果。")
    doc.add_paragraph("图1 不同受试者的准确率对比（n=42）")
    doc.save(source)

    document = parse_document(source, source.name, None)
    document.meta["image_count"] = 1

    assert "sample_size_conflict" not in _checks(document)


def test_sample_size_conflict_detected(tmp_path: Path) -> None:
    """题注 n=12 与正文「42 例」矛盾时报出，并挂在引用该图的正文段上。"""
    source = tmp_path / "conflict.docx"
    doc = Document()
    doc.add_heading("3. 结果", level=1)
    doc.add_paragraph("如图1所示，本研究共纳入42例受试者的准确率结果。")
    doc.add_paragraph("图1 不同受试者的准确率对比（n=12）")
    doc.save(source)

    document = parse_document(source, source.name, None)
    document.meta["image_count"] = 1

    conflict = [f for f in check_figure_consistency(document) if f.provenance.get("figure_check") == "sample_size_conflict"]
    assert conflict, "题注 n=12 与正文 42 例矛盾应被检出"
    assert conflict[0].checklist_item_id == "R-FIG-01"
    assert conflict[0].anchors, "冲突结论应挂在引用该图的正文段上"
