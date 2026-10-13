"""回归测试：锁定已修复的解析/判定/契约缺陷，防止再次退化。

每条用例对应一个真实缺陷，注释里写明「修之前会怎样」，避免后来者把修复当成冗余代码删掉。
"""
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from datetime import datetime, timedelta, timezone

from app.api.routes_chat import _build_question
from app.api.routes_compare import _build_response, _detail
from app.core.config import Settings, set_settings_override
from app.core.errors import AppError
from app.core.logging import key_fingerprint
from app.engine.checklist import ChecklistError, load_checklist
from app.engine.elements import extract_elements
from app.engine.figure_consistency import check_figure_consistency, find_captions
from app.engine.figures import demote_for_figures, find_figure_references
from app.engine.numeric_consistency import check_numeric_consistency
from app.engine.prompts import MAX_PROMPT_CHARS, limit_prompt_text, max_prompt_chars
from app.engine.reviewer import _apply_gates
from app.engine.revision import _section_paragraph_index, build_revised_document
from app.engine.verification import verify_findings
from app.engine.vision import read_figure_values
from app.models.schemas import (
    Anchor,
    Certainty,
    DocumentIR,
    Element,
    Finding,
    ImageRef,
    Paragraph,
    ReviewReport,
    Section,
    Severity,
    Table,
    Verdict,
)
from app.parsers.docx_parser import _image_caption, is_citation_span
from app.parsers.registry import parse_document
from app.reports.exporter import MarkdownExporter
from app.storage.db import connect, init_db
from app.storage.maintenance import cleanup_expired
from app.storage.repository import Store
from app.tasks.executor import _run_job

ROOT = Path(__file__).resolve().parents[2]


def _doc(texts: list[str], meta: dict | None = None) -> DocumentIR:
    return DocumentIR(
        source_name="t.docx",
        parser="docx",
        parser_version="t",
        paragraphs=[Paragraph(index=index, text=text) for index, text in enumerate(texts)],
        meta=meta or {},
    )


def _ids(findings: list[Finding]) -> list[str]:
    return [finding.checklist_item_id for finding in findings]


# --------------------------------------------------------------- 解析器

def test_bare_superscript_after_punctuation_does_not_crash(tmp_path: Path) -> None:
    """修前：`MATH_PREV + "%"` 是 set + str，上标裸数字跟在句号后必抛 TypeError → 上传 500。"""
    doc = Document()
    doc.add_paragraph("近年来深度学习方法发展迅速。")
    paragraph = doc.add_paragraph("本文方法取得了良好效果。")
    cite = paragraph.add_run("1")
    cite.font.superscript = True
    paragraph.add_run("随后我们给出了详细分析。")
    source = tmp_path / "bare_citation.docx"
    doc.save(source)

    document = parse_document(source, source.name, None)
    assert len(document.citations) == 1
    assert document.citations[0].numbers == [1]

    assert is_citation_span("1", "", "x") is True
    assert is_citation_span("1", "", "%") is False, "「1%」是百分比不是引用角标"


def test_multi_bracket_superscript_yields_each_citation(tmp_path: Path) -> None:
    """修前：整体只接受单个方括号组，「[1][2]」被拒 → 两条引用一起丢失。"""
    doc = Document()
    paragraph = doc.add_paragraph("已有研究表明")
    cite = paragraph.add_run("[1][2]")
    cite.font.superscript = True
    paragraph.add_run("该结论成立。")
    source = tmp_path / "multi_bracket.docx"
    doc.save(source)

    document = parse_document(source, source.name, None)
    assert [citation.text for citation in document.citations] == ["[1]", "[2]"]


def test_numbered_reference_heading_is_recognised(tmp_path: Path) -> None:
    """修前：标题必须与 REFERENCE_TITLES 完全相等，「7 References」「6 参考文献」都不认。"""
    doc = Document()
    paragraph = doc.add_paragraph("已有研究表明该方法有效")
    cite = paragraph.add_run("[1]")
    cite.font.superscript = True
    doc.add_heading("7 References", level=1)
    doc.add_paragraph("[1] Zhang S. A paper[J]. Journal, 2020.")
    source = tmp_path / "numbered_refs.docx"
    doc.save(source)

    document = parse_document(source, source.name, None)
    assert len(document.references) == 1
    assert not document.citations[0].unresolved, "参考文献已识别，不应报悬空引用"


def test_reference_entry_without_space_after_number(tmp_path: Path) -> None:
    """修前：条目正则要求编号后必须有空白，「[1]张三…」（GB/T 7714、CNKI 导出）全部漏解析。"""
    doc = Document()
    doc.add_heading("参考文献", level=1)
    doc.add_paragraph("[1]张三,李四.论文题名[J].期刊,2020.")
    doc.add_paragraph("[2]王五.另一篇[J].期刊,2021.")
    source = tmp_path / "cnki_refs.docx"
    doc.save(source)

    document = parse_document(source, source.name, None)
    assert {entry.number for entry in document.references} == {1, 2}


def test_hyperlink_text_is_not_lost(tmp_path: Path) -> None:
    """修前：只遍历 w:p 的直接子节点，w:hyperlink 内的 run 整段丢失（纯链接段落直接消失）。"""
    doc = Document()
    paragraph = doc.add_paragraph("见链接：")
    rel_id = paragraph.part.relate_to("https://example.com/paper", RT.HYPERLINK, is_external=True)
    hyperlink = paragraph._p.makeelement(qn("w:hyperlink"), {qn("r:id"): rel_id})
    hyperlink.append(paragraph.add_run("https://example.com/paper")._r)
    paragraph._p.append(hyperlink)
    doc.add_paragraph("普通段落。")
    source = tmp_path / "hyperlink.docx"
    doc.save(source)

    document = parse_document(source, source.name, None)
    assert any("example.com" in paragraph.text for paragraph in document.paragraphs)


def test_vertical_merge_cell_is_masked(tmp_path: Path) -> None:
    """修前：用 `cell is previous_row[column]` 比较，python-docx 每次新建包装对象 → 纵合并内容重复。"""
    doc = Document()
    doc.add_paragraph("表1 主结果")
    table = doc.add_table(rows=3, cols=3)
    for row in range(3):
        for column in range(3):
            table.cell(row, column).text = f"{row}{column}"
    table.cell(1, 0).merge(table.cell(2, 0))
    source = tmp_path / "vmerge.docx"
    doc.save(source)

    document = parse_document(source, source.name, None)
    assert document.tables[0].rows[2][0] is None, "纵向合并的延续单元格应以 None 占位"


def test_tab_and_formula_paragraph_is_not_a_heading(tmp_path: Path) -> None:
    """修前：all() 作用于空序列恒为 True，只含制表符+公式的段落被误判为标题。"""
    doc = Document()
    doc.add_paragraph("正文说明。")
    paragraph = doc.add_paragraph()
    paragraph.add_run().add_tab()
    paragraph._p.append(paragraph._p.makeelement(qn("m:oMath"), {}))
    doc.add_paragraph("后续正文。")
    source = tmp_path / "formula.docx"
    doc.save(source)

    document = parse_document(source, source.name, None)
    assert not [paragraph for paragraph in document.paragraphs if paragraph.is_heading]


# --------------------------------------------------------------- 图表一致性

def test_word_containing_figure_char_is_not_a_reference() -> None:
    """修前：「本研究代表3个模型」被当成「表3」引用 → 幻影 R-FIG-01 结论。"""
    document = _doc(["本研究代表3个模型的集成结果。", "表1 主结果"], meta={"image_count": 1})
    document.tables.append(
        Table(id="t1", index=0, caption="表1 主结果", rows=[["a", "b"]], n_rows=1, n_cols=2, anchor=Anchor(paragraph_index=1))
    )

    checks = {finding.provenance.get("figure_check") for finding in check_figure_consistency(document)}
    assert "reference_without_caption" not in checks
    assert all(reference["text"] != "表3" for reference in find_figure_references(document))


def test_body_sentence_does_not_shadow_caption() -> None:
    """修前：整句陈述「图 1 显示了…。」被判为题注并 setdefault 抢占，真实题注被遮蔽，
    题注与正文的样本量冲突因此永远检不出来。"""
    document = _doc(
        [
            "本研究代表3个模型的集成结果。",
            "图 1 显示了不同模型在测试集上的准确率对比。",
            "补充说明。",
            "图1 准确率对比（n=42）",
            "如图 1 所示，共 30 例样本参与评估。",
        ],
        meta={"image_count": 1},
    )

    assert find_captions(document).get(("图", 1)) == 3
    checks = {finding.provenance.get("figure_check") for finding in check_figure_consistency(document)}
    assert "sample_size_conflict" in checks


# --------------------------------------------------------------- 数值自洽

def test_unrelated_percentages_are_not_summed() -> None:
    """修前：邻近窗口内所有百分比一并求和，「80% + 20% + 正类占比 5%」被报成 105%（确定性结论）。"""
    assert check_numeric_consistency(_doc(["训练集占 80%，验证集占 20%，其中训练集正类占比 5%。"])) == []


def test_split_share_percentages_still_checked() -> None:
    """分组不得放过真正的划分合计错误。"""
    assert "R-DATA-02" in _ids(check_numeric_consistency(_doc(["按 70%、20%、5% 划分训练、验证与测试集。"])))
    assert "R-DATA-02" in _ids(
        check_numeric_consistency(_doc(["训练集 70%，验证集 20%，测试集 5%，共 100 例。"]))
    )


def test_percent_formatted_confidence_interval_is_checked() -> None:
    """修前：区间正则不接受 %，中文论文常见的「97%（95% CI 88%–96%）」完全不检查。"""
    assert "R-DATA-03" in _ids(check_numeric_consistency(_doc(["准确率 97%（95% CI 88%–96%）。"])))
    assert check_numeric_consistency(_doc(["准确率 92%（95% CI 88%–96%）。"])) == []
    # 单位混用时口径不明，保守跳过而不是误报
    assert check_numeric_consistency(_doc(["准确率 92%（95% CI 0.88–0.96）。"])) == []


# --------------------------------------------------------------- 门控与质量卡

def test_table_origin_marker_drives_demotion() -> None:
    """修前：只看锚点段落号，表格上方正文段的证据被误当「表内证据」→ 该降级的结论未降级。"""
    finding = Finding(
        finding_id="F1",
        checklist_item_id="R-12",
        verdict=Verdict.PROBLEM,
        severity=Severity.MID,
        headline="仅给出超参数取值",
        provenance={"engine": "rule", "rule_id": "R-12"},
    )
    document = _doc(["正文段", "占位段"])
    document.tables.append(
        Table(id="t1", index=0, rows=[["x"]], n_rows=1, n_cols=1, anchor=Anchor(paragraph_index=7))
    )
    body_element = [Element(id="E1", type="hyperparameter", anchors=[Anchor(paragraph_index=7)])]
    table_element = [
        Element(id="E2", type="hyperparameter", anchors=[Anchor(paragraph_index=7)], in_table=True)
    ]

    assert demote_for_figures([finding], document=document, elements=body_element)[0].verdict == Verdict.UNCERTAIN
    assert demote_for_figures([finding], document=document, elements=table_element)[0].verdict == Verdict.PROBLEM


def test_quality_card_ignores_out_of_checklist_items() -> None:
    """修前：R-DATA-*/R-FIG-01 计入审查率分子，可把 17 条清单的审查率抬到 100% 以上。"""
    findings = [
        Finding(finding_id="F1", checklist_item_id="R-01", verdict=Verdict.PASS, severity=Severity.LOW, headline="x", provenance={"engine": "rule"}),
        Finding(finding_id="F2", checklist_item_id="R-02", verdict=Verdict.PASS, severity=Severity.LOW, headline="x", provenance={"engine": "rule"}),
        Finding(finding_id="F3", checklist_item_id="R-DATA-01", verdict=Verdict.PROBLEM, severity=Severity.LOW, headline="y", provenance={"engine": "rule"}),
    ]
    _, quality = _apply_gates(findings, applicable_items=3, checklist_ids={"R-01", "R-02", "R-03"})
    assert quality["assessed_items"] == 2
    assert quality["applicable_items"] == 3
    assert quality["review_rate"] == round(2 / 3, 4)


# --------------------------------------------------------------- 契约

def test_compare_payload_carries_revision_status() -> None:
    """修前：revision_status 只拼进 notes 散文，前端读不到字段 → 每条都显示「未自动改写」。"""
    def report(report_id: str, item_id: str) -> ReviewReport:
        return ReviewReport(
            report_id=report_id,
            document_name="t.docx",
            checklist_version="1.0",
            findings=[
                Finding(
                    finding_id=f"F-{report_id}",
                    checklist_item_id=item_id,
                    verdict=Verdict.PROBLEM,
                    severity=Severity.MID,
                    certainty=Certainty.SUSPECTED,
                    headline="缺少信息",
                    provenance={"engine": "rule"},
                )
            ],
            counts={"total": 1},
            duration_ms=1,
        )

    payload = _build_response(
        report("rep-before", "R-09"),
        report("rep-after", "R-09"),
        revision_status={"R-09": ("manual", "依赖原始数据")},
        document_id="doc-revised",
    )

    assert payload["after_findings"][0]["revision_status"] == "manual"
    assert payload["after_findings"][0]["revision_note"] == "依赖原始数据"
    assert payload["kept_findings"][0]["revision_status"] == "manual"
    assert payload["document_id"] == "doc-revised"

    # 未提供状态映射时（比较两份既有报告）字段仍在，取缺省值而不是缺键
    plain = _detail(report("rep-x", "R-09").findings[0])
    assert plain["revision_status"] == "pending"
    assert plain["revision_note"] == ""


def test_checklist_rejects_non_mapping_entry(tmp_path: Path) -> None:
    """修前：条目不是对象时 `ReviewItem(**raw)` 抛 TypeError，逃出 ChecklistError 契约。"""
    path = tmp_path / "bad.json"
    path.write_text('{"items": [123]}', encoding="utf-8")
    try:
        load_checklist(path)
    except ChecklistError:
        return
    raise AssertionError("非对象条目未被转化为 ChecklistError")


def test_checklist_section_flags_heuristic_headings(tmp_path: Path) -> None:
    """启发式（非 Word 样式）标题应标记待人工确认，供前端提示章节边界不可靠。"""
    doc = Document()
    heading = doc.add_paragraph()
    heading.add_run("2. Methods").bold = True
    doc.add_paragraph("We trained the model with a learning rate of 0.001.")
    source = tmp_path / "heuristic.docx"
    doc.save(source)

    document = parse_document(source, source.name, None)
    assert document.sections
    assert all(section.needs_review for section in document.sections)


def test_elements_from_tables_are_marked(tmp_path: Path) -> None:
    """表体要素必须带 in_table 标记，供图表降级判定区分表内证据与正文证据。"""
    doc = Document()
    doc.add_paragraph("As shown in Figure 1, the cohort is summarised below.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Cohort"
    table.cell(0, 1).text = "Size"
    table.cell(1, 0).text = "训练队列"
    table.cell(1, 1).text = "42 例"
    source = tmp_path / "table_origin.docx"
    doc.save(source)

    document = parse_document(source, source.name, None)
    elements = extract_elements(document)
    assert any(element.type == "sample_size" and element.in_table for element in elements)
    assert any(not element.in_table for element in elements), "正文段要素不应被标记为表内"


# --------------------------------------------------------------- 模型输出健壮性

class _RawContentProvider:
    """返回给定原始字符串的假 provider（用于构造畸形模型输出）。"""

    def __init__(self, content: str) -> None:
        self.content = content

    def complete_structured(self, messages, temperature: float = 0.0, max_tokens=None):
        from app.engine.llm_provider import LLMResult

        return LLMResult(content=self.content, model="fake", tokens={"input": 1, "output": 1}, retries=0, duration_ms=1)


def test_quote_verification_tolerates_non_object_entries() -> None:
    """修前：results 里出现字符串元素时 `entry.get` 抛 AttributeError，逃出「解析失败即降级」契约，
    把一个本该照常出报告的审查变成失败任务。"""
    finding = Finding(
        finding_id="F1",
        checklist_item_id="R-09",
        verdict=Verdict.PROBLEM,
        severity=Severity.MID,
        headline="缺少效应量或置信区间",
        description="报告了指标但未给出置信区间。",
        provenance={"engine": "llm"},
    )
    document = _doc(["准确率为 0.91（n=42）。"])
    provider = _RawContentProvider('{"results": ["oops", 123, null, {"index": "x"}]}')

    returned = verify_findings(document, [finding], provider, max_items=6)
    assert returned == [finding], "畸形模型输出应原样返回结论，不得抛异常"


def test_vision_read_degrades_on_unexpected_worker_error(tmp_path: Path) -> None:
    """修前：并发识读路径下工作线程的非 LLMUnavailable 异常会让整次审查失败，
    与 vision.py 声明的「优雅降级、不影响主流程」不符。"""
    media_root = tmp_path / "media"
    media_root.mkdir()
    (media_root / "fig1.png").write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 32)

    document = DocumentIR(
        source_name="t.docx",
        parser="docx",
        parser_version="t",
        paragraphs=[Paragraph(index=0, text="正文")],
        images=[
            ImageRef(
                id="img_1",
                index=0,
                media_type="image/png",
                byte_size=40,
                filename="fig1.png",
                caption="图1 准确率",
                anchor=Anchor(paragraph_index=0),
            )
        ],
    )

    class _BoomProvider:
        def complete_structured(self, messages, temperature: float = 0.0, max_tokens=None):
            raise RuntimeError("boom")

    findings, tokens = read_figure_values(document, _BoomProvider(), media_root, max_images=4)
    assert findings == []
    assert tokens == {"input": 0, "output": 0}


# --------------------------------------------------------------- 追问作用域

def _report_with_finding(finding_id: str) -> ReviewReport:
    return ReviewReport(
        report_id="rep_1",
        document_name="t.docx",
        checklist_version="1.0",
        findings=[
            Finding(
                finding_id=finding_id,
                checklist_item_id="R-09",
                verdict=Verdict.PROBLEM,
                severity=Severity.MID,
                headline="缺少置信区间",
                description="未给出 95% 置信区间。",
                anchors=[Anchor(paragraph_index=3)],
                provenance={"engine": "rule"},
            )
        ],
        counts={"total": 1},
        duration_ms=1,
    )


def test_unknown_finding_id_does_not_become_empty_fulltext_qa() -> None:
    """修前：finding_id 查不到时静默回落到全文分支，而调用方在 finding 模式下并未加载文档，
    于是模型收到空 DATA 区块，却被告知「已授权通读全文」。"""
    report = _report_with_finding("F-known")

    question, fulltext = _build_question(report, "F-known", "为什么算问题？", "")
    assert fulltext is False
    assert "缺少置信区间" in question

    try:
        _build_question(report, "F-stale", "为什么算问题？", "")
    except AppError as exc:
        assert exc.code == "finding_not_found"
    else:
        raise AssertionError("失效 finding_id 未被拦截")

    question, fulltext = _build_question(report, None, "全文怎么看？", "论文全文内容")
    assert fulltext is True
    assert "论文全文内容" in question


# --------------------------------------------------------------- 修改稿生成

def test_out_of_checklist_findings_are_not_auto_rewritten() -> None:
    """修前：R-DATA-* 的 suggestion（「核对…使其之和为 10」这类自检提示）会被当成论文正文
    追加进修改稿，并把该条谎报为「已自动改写」。"""
    findings = [
        Finding(
            finding_id="F1",
            checklist_item_id="R-DATA-01",
            verdict=Verdict.PROBLEM,
            severity=Severity.LOW,
            certainty=Certainty.DETERMINISTIC,
            headline="数据划分比例之和自洽性存疑",
            description="划分比例 7:1:5 之和为 13。",
            suggestion="核对训练/验证/测试集划分比例，使其之和为 10（份数）或 100（百分比）。",
            anchors=[Anchor(paragraph_index=0)],
            provenance={"engine": "rule"},
        )
    ]
    document = _doc(["数据按 7:1:5 划分训练、验证与测试集。"])

    revised, status = build_revised_document(document, findings)

    assert status["R-DATA-01"][0] == "manual"
    assert revised.paragraphs[0].text == document.paragraphs[0].text, "不得把自检提示补写进正文"


# --------------------------------------------------- 比例口径 / 图表引用 / 题注

def test_image_caption_searches_below_then_above() -> None:
    """题注识别此前只向下找 3 段；题注位于图片上方（部分英文期刊版式）时完全漏识别。
    现改为下方优先、再向上找，且两向都要求匹配题注版式，避免把正文提及图号的句子当题注。"""
    below = [Paragraph(index=0, text="正文说明。"), Paragraph(index=1, text="图1 下方题注")]
    assert _image_caption(below, 0) == ("图1 下方题注", 1)

    above = [Paragraph(index=0, text="图2 上方题注"), Paragraph(index=1, text="正文说明。")]
    assert _image_caption(above, 1) == ("图2 上方题注", 0)

    # 下方优先：两向都有候选时取下方
    both = [Paragraph(index=0, text="图4 上方"), Paragraph(index=2, text="图4 下方")]
    assert _image_caption(both, 1) == ("图4 下方", 2)

    # 正文里提及图号的句子不是题注（不以 图N 开头）
    prose = [Paragraph(index=0, text="如图3所示，模型表现更好。"), Paragraph(index=1, text="正文说明。")]
    assert _image_caption(prose, 1) == ("", None)

    # 窗口之外不认
    far = [Paragraph(index=0, text="图5 太远了"), Paragraph(index=9, text="正文说明。")]
    assert _image_caption(far, 9) == ("", None)

def test_valid_ratio_notations_are_not_flagged() -> None:
    """修前：只认「之和为 10 或 100」，于是有效写法「4:1:1」被判为**确定性**不一致；
    且 2 段写法（「8:2 划分」）完全不检查。"""
    assert check_numeric_consistency(_doc(["训练集:验证集:测试集 = 4:1:1 划分。"])) == []
    assert check_numeric_consistency(_doc(["训练集与测试集按 8:2 划分，随机种子固定为 42。"])) == []

    # 新增覆盖：2 段写法之和既非 10 也非 100 时应当报出
    assert "R-DATA-01" in _ids(check_numeric_consistency(_doc(["训练集与测试集按 8:3 划分。"])))
    # 既有口径不回归
    assert "R-DATA-01" in _ids(check_numeric_consistency(_doc(["数据按 7:1:5 划分训练、验证与测试集。"])))
    assert "R-DATA-01" in _ids(check_numeric_consistency(_doc(["按 70:15:30 划分训练、验证与测试集。"])))


def test_figure_refs_exclude_ordinary_words() -> None:
    """修前：单字匹配「图/表」，「地图3」「试图3」「发表3篇」都会产生幻影图表引用。"""
    document = _doc(
        [
            "我们在第 3 节给出了地图3的可视化方案。",
            "作者试图3次修正标注口径。",
            "近五年发表3篇同类研究。",
            "参见图1与表2的结果。",
            "插图1展示了整体框架。",
        ]
    )
    by_paragraph: dict[int, set[str]] = {}
    for reference in find_figure_references(document):
        by_paragraph.setdefault(reference["paragraph_index"], set()).add(reference["text"].replace(" ", ""))

    assert not (by_paragraph.keys() & {0, 1, 2}), f"普通词语被误判为图表引用：{by_paragraph}"
    assert by_paragraph.get(3) == {"图1", "表2"}
    assert by_paragraph.get(4) == {"图1"}, "「插图1」应被识别为图表引用"


def test_caption_without_separator_is_recognised() -> None:
    """修前：题注必须紧跟分隔符，无分隔符写法（「图1准确率对比」）漏识别；
    同时「图1报告了…」这类正文句不得被当成题注。"""
    document = _doc(["图1准确率对比", "图1报告了不同模型的准确率。"], meta={"image_count": 1})
    assert find_captions(document).get(("图", 1)) == 0

    body_only = _doc(["图1报告了不同模型的准确率。"], meta={"image_count": 1})
    assert find_captions(body_only) == {}, "正文句不得被当成题注"


def test_revision_target_stays_inside_matched_section() -> None:
    """修前：匹配到的节若没有自己的正文段，补写会越界写到后一节的正文里。"""
    document = DocumentIR(
        source_name="t.docx",
        parser="docx",
        parser_version="t",
        paragraphs=[
            Paragraph(index=0, text="讨论", is_heading=True),
            Paragraph(index=1, text="结论", is_heading=True),
            Paragraph(index=2, text="结论正文。"),
        ],
        sections=[
            Section(id="s1", title="讨论", paragraph_index=0),
            Section(id="s2", title="结论", paragraph_index=1),
        ],
    )
    assert _section_paragraph_index(document, r"(讨论|局限|结论)") is None

    inside = DocumentIR(
        source_name="t.docx",
        parser="docx",
        parser_version="t",
        paragraphs=[
            Paragraph(index=0, text="讨论", is_heading=True),
            Paragraph(index=1, text="讨论正文。"),
            Paragraph(index=2, text="结论", is_heading=True),
            Paragraph(index=3, text="结论正文。"),
        ],
        sections=[
            Section(id="s1", title="讨论", paragraph_index=0),
            Section(id="s2", title="结论", paragraph_index=2),
        ],
    )
    assert _section_paragraph_index(inside, r"(讨论|局限|结论)") == 1


def test_prompt_budget_follows_settings() -> None:
    """提示词预算此前是硬编码常量；现由配置统一驱动，便于按模型上下文窗口调整。"""
    assert limit_prompt_text("短文本") == "短文本"
    assert max_prompt_chars() == MAX_PROMPT_CHARS

    set_settings_override(Settings(prompt_max_chars=20))
    try:
        assert max_prompt_chars() == 20
        limited = limit_prompt_text("x" * 50)
        assert limited.startswith("x" * 20)
        assert "已截断" in limited
    finally:
        set_settings_override(None)
    assert max_prompt_chars() == MAX_PROMPT_CHARS


def test_exporter_does_not_repeat_uncertain_findings() -> None:
    """修前：存疑结论在主清单与「存疑」小节各列一次，导出的 Markdown 整段重复。"""
    report = ReviewReport(
        report_id="rep_1",
        document_name="t.docx",
        checklist_version="1.0",
        findings=[
            Finding(
                finding_id="F1",
                checklist_item_id="R-09",
                verdict=Verdict.UNCERTAIN,
                severity=Severity.MID,
                certainty=Certainty.SUSPECTED,
                headline="缺少效应量或置信区间",
                description="报告了指标但未给出置信区间。",
                provenance={"engine": "rule"},
            )
        ],
        counts={"uncertain": 1},
        duration_ms=1,
    )
    text = MarkdownExporter().render(report)
    assert text.count("缺少效应量或置信区间") == 1, "存疑结论不应在导出中重复出现"
    assert "存疑（需人工复核）1 项" in text


# --------------------------------------------------- 密钥指纹 / 语种 / 任务收尾

def test_key_fingerprint_does_not_expose_key_material() -> None:
    """修前：日志直接记密钥尾 6 位——尾字符同属密钥材料，落盘即是一次泄漏面。"""
    secret = "sk-abcdefghijklmnopqrstuvwxyz012345"
    fingerprint = key_fingerprint(secret)

    assert fingerprint == key_fingerprint(secret), "同一把 Key 指纹必须恒定（可跨日志比对）"
    assert fingerprint != key_fingerprint(secret + "x"), "不同 Key 指纹必须不同"
    # 指纹是定长 hex 摘要：不含密钥本体，也不含密钥尾串（单字符偶合属摘要正常现象，不作断言）
    assert secret not in fingerprint
    assert secret[-6:] not in fingerprint
    assert len(fingerprint) == 8
    assert key_fingerprint("") == "-"


def test_document_language_is_detected(tmp_path: Path) -> None:
    """修前：`DocumentIR.language` 从未被赋值，恒为默认值，下游按语种分支时会一直走中文。"""
    chinese = Document()
    chinese.add_paragraph("本文提出一种改进的注意力机制模块，并在测试集上评估其性能。")
    zh_path = tmp_path / "zh.docx"
    chinese.save(zh_path)
    assert parse_document(zh_path, zh_path.name, None).language == "zh"

    english = Document()
    english.add_paragraph(
        "We propose an improved attention module and evaluate its performance on the test set."
    )
    en_path = tmp_path / "en.docx"
    english.save(en_path)
    assert parse_document(en_path, en_path.name, None).language == "en"


def test_job_finalization_failure_does_not_escape_worker() -> None:
    """修前：`mark_job_running` 在 try 之外、终态写入未设防，写库异常会逃出工作线程，
    任务永远停在 queued/running。"""
    class _BrokenStore:
        def __init__(self) -> None:
            self.states: list[str] = []

        def mark_job_running(self, job_id: str) -> None:
            raise RuntimeError("db locked")

        def finish_job(self, job_id: str, state: str, body: dict) -> None:
            self.states.append(state)
            raise RuntimeError("db locked")

    store = _BrokenStore()
    _run_job(store, "job_1", lambda: {"ok": True})  # 不得抛出
    assert store.states == ["failed"], "mark_job_running 失败时也应收敛到终态"

    class _OnlyFinishBroken(_BrokenStore):
        def mark_job_running(self, job_id: str) -> None:
            return None

    healthy_but_broken_finish = _OnlyFinishBroken()
    _run_job(healthy_but_broken_finish, "job_2", lambda: {"ok": True})  # 不得抛出
    assert healthy_but_broken_finish.states == ["succeeded"]


def test_cleanup_keeps_sessions_with_running_jobs(tmp_path: Path) -> None:
    """修前：过期会话会被直接删除，即使它还有 running 任务——任务收尾时结果会写进已删除的
    job（对前端等于静默丢失）并留下孤儿报告行。"""
    db_path = str(tmp_path / "app.db")
    temp_dir = tmp_path / "tmp"
    init_db(db_path)
    store = Store(db_path)
    session_id, _ = store.create_session(ttl_hours=1)
    store.create_job("job_running", session_id, "review", {})
    store.mark_job_running("job_running")

    old = (datetime.now(timezone.utc) - timedelta(hours=5)).isoformat()
    connection = connect(db_path)
    try:
        connection.execute(
            "UPDATE sessions SET created_at = ?, expires_at = ? WHERE id = ?", (old, old, session_id)
        )
        connection.commit()
    finally:
        connection.close()

    first = cleanup_expired(db_path, str(temp_dir), retention_hours=2)
    assert first["removed_sessions"] == 0, "仍有 running 任务时不得清理其会话"
    connection = connect(db_path)
    try:
        assert connection.execute("SELECT 1 FROM sessions WHERE id = ?", (session_id,)).fetchone()
    finally:
        connection.close()

    # 任务收敛到终态后，下一轮即可回收该会话
    store.finish_job("job_running", "succeeded", {"result": {}})
    second = cleanup_expired(db_path, str(temp_dir), retention_hours=2)
    assert second["removed_sessions"] == 1



