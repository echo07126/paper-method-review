from pathlib import Path

from docx import Document

from app.engine.checklist import load_checklist
from app.engine.elements import extract_elements, has_element
from app.engine.reviewer import review_document
from app.parsers.registry import parse_document

ROOT = Path(__file__).resolve().parents[2]

EN_TEXT = [
    "We evaluated the model on 1,240 fundus images, split into training and test sets.",
    "A paired t-test showed a significant difference (p < 0.05) with a 95% confidence interval.",
    "We compared against the baseline ResNet-50 and set the random seed to 42.",
    "We conducted an ablation study on the attention module.",
    "Limitations of this study are discussed in the final section.",
]


def build(path: Path, use_styles: bool, paragraphs: list[str]) -> None:
    doc = Document()
    if use_styles:
        doc.add_heading("2. Methods", level=1)
    else:
        heading = doc.add_paragraph()
        heading.add_run("2. Methods").bold = True
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(path)


def test_english_paper_no_false_positives(tmp_path: Path) -> None:
    source = tmp_path / "english.docx"
    build(source, True, EN_TEXT)
    document = parse_document(source, source.name, None)
    elements = extract_elements(document)
    assert has_element(elements, "sample_size"), "英文样本量应被识别"
    assert has_element(elements, "stat_test"), "t-test 应被识别为统计检验"
    assert has_element(elements, "limitation"), "Limitations 应被识别"

    report = review_document(document, load_checklist(ROOT / "checklists" / "quant-ai-v1.json"))
    problems = {f.checklist_item_id for f in report.findings if f.verdict.value == "problem"}
    assert "R-07" not in problems, "已报告 t 检验时不应误报 R-07"
    assert "R-15" not in problems, "已讨论局限时不应误报 R-15"


def test_bold_only_headings_detected(tmp_path: Path) -> None:
    source = tmp_path / "bold_heading.docx"
    build(source, False, EN_TEXT)
    document = parse_document(source, source.name, None)
    assert document.sections, "无 Word 样式时也应通过启发式识别标题"
    assert document.meta.get("heuristic_headings", 0) >= 1


def test_table_text_is_reviewed(tmp_path: Path) -> None:
    source = tmp_path / "with_table.docx"
    doc = Document()
    doc.add_heading("3. Results", level=1)
    doc.add_paragraph("We compared the proposed model with the baseline method.")
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "Model"
    table.cell(0, 1).text = "Accuracy"
    table.cell(0, 2).text = "p value"
    table.cell(1, 0).text = "Ours"
    table.cell(1, 1).text = "0.91"
    table.cell(1, 2).text = "p<0.05"
    doc.save(source)

    document = parse_document(source, source.name, None)
    assert document.meta.get("table_count") == 1
    assert len(document.tables) == 1
    table = document.tables[0]
    assert table.n_rows == 2 and table.n_cols == 3
    assert table.header_rows == 1, "首行无数字，应识别为表头"
    assert table.rows[1][1] == "0.91"
    elements = extract_elements(document)
    assert has_element(elements, "p_value"), "表格中的 p 值应被抽取"
    assert all(paragraph.style != "Table" for paragraph in document.paragraphs), "表格应移出段落流"
