import json
from pathlib import Path

from docx import Document

from app.engine.checklist import ChecklistError, load_checklist
from app.engine.elements import extract_elements, has_element
from app.engine.reviewer import review_document
from app.engine.rules import run_rules
from app.parsers.registry import parse_document

ROOT = Path(__file__).resolve().parents[2]


def build_methods_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("2. 方法", level=1)
    doc.add_paragraph("本研究纳入 1,240 例眼底彩照，按 8:2 划分为训练集与验证集。")
    doc.add_paragraph("超参数（学习率、批大小）在验证集上通过网格搜索确定，最终在测试集上评估。")
    doc.add_paragraph("结果显示准确率为 0.912，较基线提升 2.1 个百分点（p<0.05）。")
    doc.add_paragraph("本文提出了一种改进的注意力机制模块，可推广至各类医疗影像场景。")
    doc.save(path)


def test_element_extraction(tmp_path: Path) -> None:
    source = tmp_path / "methods.docx"
    build_methods_docx(source)
    document = parse_document(source, "methods.docx", None)
    elements = extract_elements(document)
    assert has_element(elements, "sample_size")
    assert has_element(elements, "data_split")
    assert has_element(elements, "hyperparameter")
    assert has_element(elements, "p_value")
    assert has_element(elements, "overgeneralization")
    assert has_element(elements, "novel_module")


def test_element_based_rules(tmp_path: Path) -> None:
    source = tmp_path / "methods.docx"
    build_methods_docx(source)
    document = parse_document(source, "methods.docx", None)
    ids = {finding.checklist_item_id for finding in run_rules(document)}
    assert "R-03" in ids, "应检出调参/泄漏风险"
    assert "R-07" in ids, "应检出 p 值但无检验方法"
    assert "R-09" in ids, "应检出缺少置信区间"
    assert "R-11" in ids, "应检出结论外推"
    assert "R-05" in ids, "应检出缺少消融"


def test_checklist_rejects_duplicate_ids(tmp_path: Path) -> None:
    payload = {"items": [
        {"id": "R-X", "category": "t", "title": "a"},
        {"id": "R-X", "category": "t", "title": "b"},
    ]}
    path = tmp_path / "bad.json"
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    try:
        load_checklist(path)
        raise AssertionError("重复 ID 未被拦截")
    except ChecklistError:
        pass


class FakeProvider:
    def complete_structured(self, messages, temperature: float = 0.0):
        from app.engine.llm_provider import LLMResult

        content = json.dumps([
            {
                "checklist_item_id": "R-05",
                "verdict": "problem",
                "severity": "mid",
                "headline": "缺少消融实验",
                "description": "模型声称缺少消融。",
                "suggestion": "补充消融。",
            }
        ])
        return LLMResult(content=content, model="fake-model", tokens={"input": 10, "output": 20}, retries=0, duration_ms=5)


def test_evidence_gate_marks_uncertain_without_anchor(tmp_path: Path) -> None:
    doc = Document()
    doc.add_heading("方法", level=1)
    doc.add_paragraph("本文仅描述了一个简单流程，没有其他信息。")
    source = tmp_path / "plain.docx"
    doc.save(source)
    document = parse_document(source, "plain.docx", None)
    items = load_checklist(ROOT / "checklists" / "quant-ai-v1.json")
    report = review_document(document, items, provider=FakeProvider(), use_llm=True)
    llm_findings = [f for f in report.findings if f.provenance.get("engine") == "llm"]
    assert llm_findings, "应包含 LLM 判定"
    assert all(f.provenance.get("evidence_gate") == "no_evidence" for f in llm_findings)
    assert all(f.verdict.value == "uncertain" for f in llm_findings)
