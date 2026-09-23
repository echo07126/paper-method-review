import json
from pathlib import Path

from docx import Document

from app.engine.checklist import load_checklist
from app.engine.llm_provider import LLMResult
from app.engine.reviewer import review_document
from app.parsers.registry import parse_document

ROOT = Path(__file__).resolve().parents[2]


def build_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("2. 方法", level=1)
    doc.add_paragraph("本文提出了一种改进的注意力机制模块，在测试集上准确率达到 0.92，训练集与测试集划分见附录。")
    doc.save(path)


class ContractProvider:
    def complete_structured(self, messages, temperature: float = 0.0):
        content = json.dumps([
            {"checklist_item_id": "R-05", "verdict": "pass", "severity": None, "headline": "认为无需消融", "description": "d", "suggestion": "s"},
            {"checklist_item_id": "R-06", "verdict": "pass", "severity": None, "headline": "指标合适", "description": "", "suggestion": ""},
            {"checklist_item_id": "R-XX", "verdict": "bad_verdict", "severity": "high", "headline": "非法", "description": "", "suggestion": ""},
        ])
        return LLMResult(content=content, model="contract-fake", tokens={"input": 5, "output": 6}, retries=0, duration_ms=1)


def test_merge_dedup_and_contract(tmp_path: Path) -> None:
    source = tmp_path / "merge.docx"
    build_docx(source)
    document = parse_document(source, "merge.docx", None)
    items = load_checklist(ROOT / "checklists" / "quant-ai-v1.json")

    report = review_document(document, items, provider=ContractProvider(), use_llm=True)

    r05 = [f for f in report.findings if f.checklist_item_id == "R-05"]
    assert len(r05) == 1, "同一清单条目应合并为一条结论"
    assert r05[0].provenance.get("merged_with_llm") is True
    assert r05[0].provenance.get("conflict") is True, "规则与模型结论不同应记录冲突"
    assert r05[0].verdict.value == "problem", "冲突时以规则结论为准"

    r06 = [f for f in report.findings if f.checklist_item_id == "R-06"]
    assert r06 and r06[0].verdict.value == "pass"
    assert report.counts["pass"] >= 1
    assert report.counts["total"] < report.counts["assessed"], "pass 不应计入问题总数"

    assert any("不符合契约" in note for note in report.notes), "非法条目应被丢弃并记录"
