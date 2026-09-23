import json
from pathlib import Path

from docx import Document

from app.engine.checklist import load_checklist
from app.engine.llm_provider import LLMResult
from app.engine.reviewer import review_document
from app.parsers.registry import parse_document

ROOT = Path(__file__).resolve().parents[2]


class QuoteProvider:
    """先给出两条无锚点的 LLM 结论，再在引用验证阶段返回一真一假引文。"""

    def complete_structured(self, messages, temperature: float = 0.0):
        system = messages[0]["content"]
        if "逐字引用" in system:
            content = json.dumps({
                "results": [
                    {"index": 0, "quote": "本模型实验仅运行一次，未报告方差与置信区间"},
                    {"index": 1, "quote": "完全不存在的句子XYZ"},
                ]
            })
        else:
            content = json.dumps([
                {"checklist_item_id": "R-10", "verdict": "problem", "severity": "mid", "headline": "缺少重复实验与方差", "description": "仅单次运行", "suggestion": "补充多次运行"},
                {"checklist_item_id": "R-15", "verdict": "problem", "severity": "low", "headline": "缺少局限", "description": "无局限讨论", "suggestion": "补充局限"},
            ])
        return LLMResult(content=content, model="quote-fake", tokens={"input": 5, "output": 6}, retries=0, duration_ms=1)


def build_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("3. 结果", level=1)
    doc.add_paragraph("本模型实验仅运行一次，未报告方差与置信区间。")
    doc.save(path)


def test_quote_verification_promotes_and_keeps_uncertain(tmp_path: Path) -> None:
    source = tmp_path / "quote.docx"
    build_docx(source)
    document = parse_document(source, "quote.docx", None)
    items = load_checklist(ROOT / "checklists" / "quant-ai-v1.json")

    report = review_document(document, items, provider=QuoteProvider(), use_llm=True)

    r10 = [f for f in report.findings if f.checklist_item_id == "R-10"][0]
    assert r10.anchors, "引文命中应挂上原文锚点"
    assert r10.provenance.get("evidence_gate") == "quote_verified"
    assert r10.verdict.value == "uncertain", "顾问模式下模型结论为存疑/建议"

    r15 = [f for f in report.findings if f.checklist_item_id == "R-15"][0]
    assert not r15.anchors
    assert r15.verdict.value == "uncertain", "引文无法匹配时应维持存疑"
    assert report.counts["uncertain"] >= 1
