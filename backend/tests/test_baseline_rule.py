from pathlib import Path

from docx import Document

from app.engine.checklist import load_checklist
from app.engine.reviewer import review_document
from app.parsers.registry import parse_document

ROOT = Path(__file__).resolve().parents[2]


def build(path: Path, paragraphs: list[str]) -> None:
    doc = Document()
    doc.add_heading("2. 方法", level=1)
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(path)


def verdict_of(report, item_id: str):
    matches = [f for f in report.findings if f.checklist_item_id == item_id]
    return matches[0].verdict.value if matches else None


def test_baseline_missing_flags_problem(tmp_path: Path) -> None:
    source = tmp_path / "no_baseline.docx"
    build(source, [
        "本研究纳入 860 例样本，用于模型训练与评估。",
        "本文提出的融合模块在多数任务上优于已有方法。",
    ])
    document = parse_document(source, source.name, None)
    report = review_document(document, load_checklist(ROOT / "checklists" / "quant-ai-v1.json"))
    assert verdict_of(report, "R-04") == "problem"


def test_baseline_present_passes(tmp_path: Path) -> None:
    source = tmp_path / "has_baseline.docx"
    build(source, [
        "本研究纳入 400 例样本并与基线方法 ResNet-50 对比。",
        "本文提出的融合模块取得更优结果。",
    ])
    document = parse_document(source, source.name, None)
    report = review_document(document, load_checklist(ROOT / "checklists" / "quant-ai-v1.json"))
    assert verdict_of(report, "R-04") == "pass"
