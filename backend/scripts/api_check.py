"""API 冒烟：健康检查 -> 上传 -> 审查 -> 报告 -> 导出 -> 会话清理（无需 pytest）。

配置隔离说明
------------
`Settings` 的取值优先级为「显式入参 > 项目 backend/.env > 系统环境变量 > secrets」，
且 `model_config.env_file` 在**类定义时**固化为 `backend/.env`。因此只设 `os.environ`
会被 `.env` 覆盖，导致冒烟写入真实 `./data`（含论文正文）、`CHECKLIST_DIR` 失效而报错。

正确做法：派生 `env_file=None` 的子类构造隔离实例，经
`app.core.config.set_settings_override()` 注入——该注入点对所有
`from app.core.config import get_settings` 的调用方统一生效。
"""
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
BACKEND = ROOT / "backend"
sys.path.insert(0, str(BACKEND))

from docx import Document  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from pydantic_settings import SettingsConfigDict  # noqa: E402

from app.core.config import Settings, set_settings_override  # noqa: E402
from app.main import app  # noqa: E402
from app.storage.db import init_db  # noqa: E402
from app.storage.repository import Store  # noqa: E402


class IsolatedSettings(Settings):
    """断开 backend/.env，使显式入参成为唯一配置来源（保证测试隔离）。"""

    model_config = SettingsConfigDict(env_file=None, extra="ignore")


def build_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("方法", level=1)
    p = doc.add_paragraph()
    p.add_run("引用了一条不存在的文献")
    sup = p.add_run("[99]")
    sup.font.superscript = True
    doc.add_heading("References", level=1)
    doc.add_paragraph("[1] Some reference. 2026.")
    doc.save(path)


def wait_task(client: TestClient, prefix: str, task_id: str, timeout: float = 30.0) -> dict:
    """轮询任务至终态，返回 result（重接口已异步化：POST 202 + GET /tasks/{id}）。"""
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        response = client.get(f"{prefix}/tasks/{task_id}")
        assert response.status_code == 200, response.text
        payload = response.json()
        if payload["state"] == "succeeded":
            return payload["result"]
        if payload["state"] == "failed":
            raise AssertionError(f"任务失败：{payload['error']}")
        time.sleep(0.05)
    raise AssertionError("任务超时未完成")


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="api_check_"))
    settings = IsolatedSettings(
        db_path=str(tmp / "app.db"),
        temp_dir=str(tmp / "tmp"),
        checklist_dir=str(ROOT / "checklists"),
        app_env="dev",
        docs_enabled=False,
    )
    set_settings_override(settings)  # 全局生效，不触碰真实 ./data
    init_db(settings.db_path)

    sample = tmp / "sample.docx"
    build_docx(sample)
    prefix = "/api/v1"

    try:
        with TestClient(app) as client:
            health = client.get(f"{prefix}/health")
            assert health.status_code == 200, health.text

            with sample.open("rb") as handle:
                upload = client.post(f"{prefix}/uploads", files={"file": ("sample.docx", handle, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
            assert upload.status_code == 200, upload.text
            document_id = upload.json()["document_id"]

            review = client.post(f"{prefix}/reviews", json={"document_id": document_id, "use_llm": False})
            assert review.status_code == 202, review.text
            review_result = wait_task(client, prefix, review.json()["task_id"])
            report_id = review_result["report_id"]

            report = client.get(f"{prefix}/reports/{report_id}")
            assert report.status_code == 200, report.text
            assert report.json()["counts"]["total"] >= 1

            exported = client.get(f"{prefix}/reports/{report_id}/export", params={"format": "markdown"})
            assert exported.status_code == 200 and "审查报告" in exported.text

            # 追问多轮上下文（需求 15.5.4）：两轮提问应带回累积历史与 session_id
            first = client.post(f"{prefix}/reports/{report_id}/chat", json={"question": "这条该怎么改？"})
            assert first.status_code == 202, first.text
            wait_task(client, prefix, first.json()["task_id"])
            second = client.post(f"{prefix}/reports/{report_id}/chat", json={"question": "有推荐的写法吗？"})
            assert second.status_code == 202, second.text
            second_result = wait_task(client, prefix, second.json()["task_id"])
            assert second_result["session_id"], "响应应返回 session_id"
            chat_session_id = second_result["session_id"]
            assert len(second_result["history"]) == 4, "两轮追问应落库 4 条消息（2 轮 × user+assistant）"

            other = TestClient(app)
            forbidden = other.get(f"{prefix}/reports/{report_id}")
            assert forbidden.status_code == 404, f"越权未拦截: {forbidden.status_code}"

            purge = client.post(f"{prefix}/sessions/purge")
            assert purge.status_code == 200

        # 删除传播：追问历史必须随会话一并清除（需求 15.5.4 / 16.2）
        assert Store(settings.db_path).list_chat_messages(chat_session_id, report_id) == [], "purge 后仍有追问历史残留"
    finally:
        set_settings_override(None)

    # 隔离自证：真实 ./data 不得被本次冒烟创建或写入
    real_data = (ROOT / "data").resolve()
    assert not real_data.exists() or not any(real_data.iterdir()), f"隔离失效：真实 data/ 被写入 {real_data}"

    print("API CHECK PASS")
    print(" - 隔离目录 :", tmp)
    print(" - health    :", health.json())
    print(" - upload    : paragraphs =", upload.json()["paragraph_count"], "warnings =", len(upload.json()["warnings"]))
    print(" - review    :", review_result["counts"])
    print(" - chat      : source =", second_result["source"], "history =", len(second_result["history"]))
    print(" - export    : length =", len(exported.text))
    print(" - 越权访问   : ->", forbidden.status_code)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
