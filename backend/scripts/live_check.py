"""对运行中的服务做真实 HTTP 端到端检查（不使用 TestClient）。

重接口已异步化：`/reviews`、`/compare`、`/reports/{id}/chat` 均**立即返回 202** 与
`task_id`，需再轮询 `GET /tasks/{task_id}` 取 `result`（与旧同步响应逐字段一致）。
"""
import io
import time

import httpx
from docx import Document

BASE = "http://127.0.0.1:8000/api/v1"


def build_docx() -> bytes:
    doc = Document()
    doc.add_heading("3. Results", level=1)
    doc.add_paragraph("Figure 1 shows the accuracy of 0.91; we compared with the baseline ResNet-50.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def wait_task(client: httpx.Client, task_id: str, timeout: float = 120.0) -> dict:
    """轮询任务至终态，返回 result。"""
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        response = client.get(f"/tasks/{task_id}")
        response.raise_for_status()
        payload = response.json()
        if payload["state"] == "succeeded":
            return payload["result"]
        if payload["state"] == "failed":
            raise AssertionError(f"任务失败：{payload['error']}")
        time.sleep(0.1)
    raise AssertionError("任务超时未完成")


def main() -> int:
    with httpx.Client(base_url=BASE, timeout=60, follow_redirects=True) as client:
        health = client.get("/health")
        print("health:", health.status_code, health.json())

        content = build_docx()
        upload = client.post("/uploads", files={"file": ("live.docx", content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        upload.raise_for_status()
        document_id = upload.json()["document_id"]
        print("upload:", upload.status_code, "sections=", len(upload.json()["sections"]), "warnings=", upload.json()["warnings"])

        strict_task = client.post("/reviews", json={"document_id": document_id, "use_llm": False, "demote_on_figures": False})
        conservative_task = client.post("/reviews", json={"document_id": document_id, "use_llm": False, "demote_on_figures": True})
        assert strict_task.status_code == 202 and conservative_task.status_code == 202, "重接口应返回 202（任务模式）"
        strict = wait_task(client, strict_task.json()["task_id"])
        conservative = wait_task(client, conservative_task.json()["task_id"])
        print("strict      :", strict["counts"], "figures=", strict["figure_references"], "demote=", strict["demote_on_figures"])
        print("conservative:", conservative["counts"], "figures=", conservative["figure_references"], "demote=", conservative["demote_on_figures"])
        assert conservative["demote_on_figures"] is True
        assert conservative["counts"]["total"] < strict["counts"]["total"]

        reports = client.get("/reports").json()
        # 基线固定取 strict（未降级）报告，避免与 conservative（已降级）报告做不对等对比
        report_id = strict["report_id"]
        print("reports:", len(reports), "-> baseline", report_id)
        report = client.get(f"/reports/{report_id}").json()
        print("report fields:", sorted(report.keys()))
        export = client.get(f"/reports/{report_id}/export", params={"format": "markdown"})
        print("export:", export.status_code, "bytes=", len(export.text))

        chat_task = client.post(f"/reports/{report_id}/chat", json={"question": "怎么改？"})
        assert chat_task.status_code == 202, "问答接口应返回 202"
        chat = wait_task(client, chat_task.json()["task_id"])
        print("chat source:", chat.get("source"))

        # 端到端：不传 after_report_id → 后端自动生成修改稿并二次审查（规则模式，零模型调用）
        compare_task = client.post("/compare", json={"before_report_id": report_id, "use_llm": False})
        assert compare_task.status_code == 202, "对比接口应返回 202"
        compare = wait_task(client, compare_task.json()["task_id"])
        print("compare: resolved=", len(compare["resolved"]), "new=", len(compare["new"]), "kept=", len(compare["kept"]))
        assert compare["after"], "对比应产出修改稿"
        assert compare["after"]["total"] <= compare["before"]["total"], "修改稿问题数应不高于原稿"

        purge = client.post("/sessions/purge")
        print("purge:", purge.status_code, purge.json())
    print("LIVE CHECK PASS")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
