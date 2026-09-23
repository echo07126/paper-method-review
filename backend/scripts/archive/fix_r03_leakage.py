"""修正 R-03：只有“在测试集上调参/做模型选择”才算泄漏；验证集调参属规范做法。

同时把 synthetic-001 夹具改成真正的泄漏场景（在测试集上调参），保持该规则的阳性用例。
"""
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
APP = ROOT / "backend" / "app"

# 1) 规则：泄漏 = 调参类词 + “测试集”，且没有“测试集仅用于最终评估”这类声明
rules = APP / "engine" / "rules.py"
text = rules.read_text(encoding="utf-8")
old_signal = 'EVAL_SET_PATTERN = re.compile(r"(验证集|测试集)")'
new_signal = (
    'TEST_SET_PATTERN = re.compile(r"(测试集|test\\s*set)", re.IGNORECASE)\n'
    'PROPER_USE_PATTERN = re.compile(r"(测试集.{0,12}(仅|只|一次性).{0,12}(用|使用|评估|评价)|test\\s*set.{0,20}(used\\s*once|only\\s*for\\s*final))", re.IGNORECASE)'
)
if "TEST_SET_PATTERN" not in text:
    text = text.replace(old_signal, new_signal)

old_block = '''    leakage: list[Finding] = []
    for paragraph in document.paragraphs:
        if LEAKAGE_PATTERN.search(paragraph.text) and EVAL_SET_PATTERN.search(paragraph.text):'''
new_block = '''    leakage: list[Finding] = []
    for paragraph in document.paragraphs:
        # 仅当“调参/模型选择”与“测试集”同时出现、且未声明“测试集仅用于最终评估”时，才判为泄漏；
        # 在验证集上调参属于规范做法，不应误报。
        if (
            LEAKAGE_PATTERN.search(paragraph.text)
            and TEST_SET_PATTERN.search(paragraph.text)
            and not PROPER_USE_PATTERN.search(paragraph.text)
        ):'''
if old_block in text:
    text = text.replace(old_block, new_block)
    print("R-03 判定逻辑已修正")
else:
    print("[warn] 未找到 R-03 判定块")
rules.write_text(text, encoding="utf-8")

# 2) 夹具：synthetic-001 改为真正的泄漏场景
from docx import Document  # noqa: E402

fixture = ROOT / "samples" / "fixtures" / "synthetic-001.docx"
doc = Document(str(fixture))
changed = False
for paragraph in doc.paragraphs:
    if "网格搜索" in paragraph.text:
        for run in paragraph.runs:
            if "网格搜索" in run.text or "验证集" in run.text:
                pass
        # 重写整段文本（保留段落格式）
        new_text = "超参数（学习率、批大小）在测试集上通过网格搜索确定，并在同一测试集上评估最终性能。"
        if paragraph.runs:
            paragraph.runs[0].text = new_text
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.add_run(new_text)
        changed = True
if changed:
    doc.save(str(fixture))
    print("synthetic-001 夹具已改为“测试集上调参”（真正的泄漏场景）")
else:
    print("[warn] 未在夹具中找到目标段落")
