"""无 pytest 依赖的冒烟测试：解析 -> 规则 -> 报告 -> 存储隔离 -> 上传校验。"""
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
BACKEND = ROOT / "backend"
sys.path.insert(0, str(BACKEND))

from docx import Document  # noqa: E402

from app.core.errors import AppError  # noqa: E402
from app.engine.checklist import load_checklist  # noqa: E402
from app.engine.reviewer import review_document  # noqa: E402
from app.parsers.registry import parse_document  # noqa: E402
from app.storage.files import save_upload  # noqa: E402
from app.storage.db import init_db  # noqa: E402
from app.storage.repository import Store  # noqa: E402


def make_sample(path: Path) -> None:
    doc = Document()
    doc.add_heading("示例论文：AI 实证方法学演示", level=0)

    def body(parts):
        paragraph = doc.add_paragraph()
        for text, sup in parts:
            run = paragraph.add_run(text)
            run.font.superscript = sup

    doc.add_heading("1. 方法", level=1)
    body([("模型在测试集上完成调参后报告准确率", False), ("[1]", True), ("，并与基线比较。", False)])
    body([("该结论引用了一条不存在的文献", False), ("[99]", True), ("。", False)])
    doc.add_heading("References", level=1)
    doc.add_paragraph("[1] Author 1. Synthetic reference for smoke test. 2026.")
    doc.save(path)


def main() -> int:
    checks: list[str] = []
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        docx_path = tmp_path / "sample.docx"
        make_sample(docx_path)

        document = parse_document(docx_path, "sample.docx", None)
        assert document.sections, "应识别出章节标题"
        assert len(document.references) == 1, "应解析出 1 条参考文献"
        dangling = [c for c in document.citations if c.unresolved]
        assert dangling and dangling[0].unresolved == [99], "应识别出悬空引用 [99]"
        checks.append(f"解析：章节 {len(document.sections)} / 引用 {len(document.citations)} / 文献 {len(document.references)}")

        items = load_checklist(ROOT / "checklists" / "quant-ai-v1.json")
        assert len(items) >= 15, "清单条目应不少于 15 条"
        report = review_document(document, items, provider=None, use_llm=False)
        assert report.counts["total"] >= 1, "规则引擎应至少产出一条问题"
        checks.append(f"审查：问题总数 {report.counts['total']}（严重 {report.counts['high']}）")

        init_db(str(tmp_path / "app.db"))
        store = Store(str(tmp_path / "app.db"))
        session_a, _ = store.create_session(2)
        document_id = store.save_document(session_a, document.source_name, document.model_dump_json())
        store.authorized_document(session_a, document_id)
        checks.append("存储：同会话读取文档成功")

        session_b, _ = store.create_session(2)
        try:
            store.authorized_document(session_b, document_id)
            raise AssertionError("越权访问未被拦截")
        except AppError as exc:
            assert exc.status == 404, "越权应返回 404"
            checks.append("安全：跨会话访问被拦截（404）")

        try:
            save_upload(str(tmp_path / "tmp"), session_a, "fake.docx", b"NOTADOCX", 50)
            raise AssertionError("伪造扩展名未被拦截")
        except AppError as exc:
            assert exc.code == "file_content_mismatch", "伪造扩展名应被识别"
            checks.append("安全：伪造扩展名被拦截")

    print("SMOKE PASS")
    for line in checks:
        print(" -", line)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
