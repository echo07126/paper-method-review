"""生成规则夹具（fixture）样本（DOCX + 标准答案 JSON）。默认不覆盖已存在文件。"""
import argparse
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend"))

from docx import Document  # noqa: E402

FIXTURES_DIR = ROOT / "samples" / "fixtures"
TRUTH_DIR = ROOT / "samples" / "ground_truth"

SAMPLES = [
    {
        "sample_id": "synthetic-001",
        "title": "眼底分类模型的调参与显著性报告（预埋：泄漏/检验/区间/种子）",
        "paragraphs": [
            ("h", "2. 方法"),
            ("p", "本研究纳入 1,240 例眼底彩照，按 8:2 划分为训练集与验证集。"),
            ("p", "超参数（学习率、批大小）在验证集上通过网格搜索确定，最终在测试集上评估。"),
            ("p", "结果显示准确率为 0.912，较基线提升 2.1 个百分点（p<0.05）。"),
            ("p", "由于资源限制，未开展额外实验。"),
        ],
        "items": [
            {"checklist_item_id": "R-03", "severity_expected": "high", "anchor": {"paragraph_index": 2}, "note": "验证集上调参"},
            {"checklist_item_id": "R-07", "severity_expected": "high", "anchor": {"paragraph_index": 3}, "note": "只有 p 值无检验方法"},
            {"checklist_item_id": "R-09", "severity_expected": "mid", "anchor": {"paragraph_index": 3}, "note": "无置信区间"},
            {"checklist_item_id": "R-02", "severity_expected": "low", "anchor": {"paragraph_index": 1}, "note": "样本量无依据"},
            {"checklist_item_id": "R-14", "severity_expected": "low", "anchor": {"paragraph_index": 1}, "note": "无随机种子"},
        ],
    },
    {
        "sample_id": "synthetic-002",
        "title": "注意力模块与结论外推（预埋：消融/外推/区间/样本量）",
        "paragraphs": [
            ("h", "2. 方法"),
            ("p", "本研究共收集 320 例样本，用于模型训练与评估。"),
            ("p", "配对 t 检验显示两组差异显著（p<0.05）。"),
            ("p", "本文提出了一种改进的注意力机制模块。"),
            ("p", "综上，所提方法普遍适用，可推广至各类任务。"),
        ],
        "items": [
            {"checklist_item_id": "R-05", "severity_expected": "mid", "anchor": {"paragraph_index": 3}, "note": "提出模块无消融"},
            {"checklist_item_id": "R-09", "severity_expected": "mid", "anchor": {"paragraph_index": 2}, "note": "有检验无区间"},
            {"checklist_item_id": "R-11", "severity_expected": "low", "anchor": {"paragraph_index": 4}, "note": "结论外推"},
            {"checklist_item_id": "R-02", "severity_expected": "low", "anchor": {"paragraph_index": 1}, "note": "样本量无依据"},
        ],
    },
    {
        "sample_id": "synthetic-003",
        "title": "引用规范与模块消融（预埋：悬空引用/缺文献表/消融）",
        "paragraphs": [
            ("h", "2. 方法"),
            ("sup", "本研究纳入 500 例样本，方法参照既往工作", "[99]", "。"),
            ("p", "组间差异采用独立样本 t 检验（p<0.05）。"),
            ("p", "本文提出的模型在多个数据集上表现良好。"),
        ],
        "items": [
            {"checklist_item_id": "R-REF-01", "severity_expected": "low", "anchor": {"paragraph_index": 1}, "note": "悬空引用 [99]"},
            {"checklist_item_id": "R-REF-02", "severity_expected": "mid", "anchor": {"paragraph_index": 1}, "note": "无参考文献章节"},
            {"checklist_item_id": "R-05", "severity_expected": "mid", "anchor": {"paragraph_index": 3}, "note": "提出模块无消融"},
            {"checklist_item_id": "R-09", "severity_expected": "mid", "anchor": {"paragraph_index": 2}, "note": "有检验无区间"},
            {"checklist_item_id": "R-02", "severity_expected": "low", "anchor": {"paragraph_index": 1}, "note": "样本量无依据"},
        ],
    },
    {
        "sample_id": "synthetic-004",
        "title": "多数据集比较（预埋：划分/检验/基线/多重比较）",
        "paragraphs": [
            ("h", "2. 方法"),
            ("p", "本研究纳入 860 例样本，用于模型训练与评估。"),
            ("p", "模型在多个数据集上进行验证，结果显示准确率为 0.88，并报告 P<0.01。"),
            ("p", "本文提出的融合模块在多数任务上优于已有方法。"),
        ],
        "items": [
            {"checklist_item_id": "R-02", "severity_expected": "low", "anchor": {"paragraph_index": 1}, "note": "样本量无依据"},
            {"checklist_item_id": "R-03", "severity_expected": "high", "anchor": {"paragraph_index": 2}, "note": "未描述数据划分"},
            {"checklist_item_id": "R-04", "severity_expected": "mid", "anchor": {"paragraph_index": 3}, "note": "未给出可核对的基线方法"},
            {"checklist_item_id": "R-05", "severity_expected": "mid", "anchor": {"paragraph_index": 3}, "note": "提出模块无消融"},
            {"checklist_item_id": "R-07", "severity_expected": "high", "anchor": {"paragraph_index": 2}, "note": "有 p 值无检验方法"},
            {"checklist_item_id": "R-08", "severity_expected": "mid", "anchor": {"paragraph_index": 2}, "note": "多数据集比较未校正"},
            {"checklist_item_id": "R-09", "severity_expected": "mid", "anchor": {"paragraph_index": 2}, "note": "无置信区间"},
            {"checklist_item_id": "R-14", "severity_expected": "low", "anchor": {"paragraph_index": 1}, "note": "无随机种子"},
            {"checklist_item_id": "R-15", "severity_expected": "low", "anchor": {"paragraph_index": 3}, "note": "无局限讨论"},
        ],
    },
    {
        "sample_id": "synthetic-005",
        "title": "规范良好样本（负对照：应无 problem 输出）",
        "paragraphs": [
            ("h", "2. 方法"),
            ("p", "本研究依据功效分析确定样本量为 400 例，并按 7:1:2 划分为训练/验证/测试集，固定随机种子 42。"),
            ("p", "组间差异采用配对 t 检验（P=0.012，95% CI 0.4–3.1），并与基线方法 ResNet-50 对比。"),
            ("p", "对提出的注意力模块进行了消融实验。"),
            ("p", "本文方法在单中心数据上初步有效；局限与未来工作见末节。"),
        ],
        "items": [],
    },
]


def build_docx(spec: dict, path: Path) -> None:
    doc = Document()
    for entry in spec["paragraphs"]:
        kind = entry[0]
        if kind == "h":
            doc.add_heading(entry[1], level=1)
        elif kind == "p":
            doc.add_paragraph(entry[1])
        elif kind == "sup":
            paragraph = doc.add_paragraph()
            paragraph.add_run(entry[1])
            marker = paragraph.add_run(entry[2])
            marker.font.superscript = True
            paragraph.add_run(entry[3])
    doc.save(path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--force", action="store_true", help="覆盖已存在的样本与标准答案")
    args = parser.parse_args()

    FIXTURES_DIR.mkdir(parents=True, exist_ok=True)
    TRUTH_DIR.mkdir(parents=True, exist_ok=True)
    created = skipped = 0
    for spec in SAMPLES:
        docx_path = FIXTURES_DIR / f"{spec['sample_id']}.docx"
        truth_path = TRUTH_DIR / f"{spec['sample_id']}.json"
        if docx_path.exists() and not args.force:
            skipped += 1
            continue
        build_docx(spec, docx_path)
        truth = {
            "sample_id": spec["sample_id"],
            "title": spec["title"],
            "type": "fixture",
            "source": None,
            "license": "self-authored",
            "language": "zh",
            "items": spec["items"],
        }
        truth_path.write_text(json.dumps(truth, ensure_ascii=False, indent=2), encoding="utf-8")
        created += 1
        print(f"generated: {docx_path.name} ({len(spec['items'])} 条标准答案)")
    print(f"完成：新增 {created} 篇，跳过已存在 {skipped} 篇")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
